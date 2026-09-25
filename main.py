import streamlit as st
import pandas as pd
import os
import re
import io
import time
import datetime
import warnings
import openpyxl
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

# ==================================================
# PAGE CONFIG
# ==================================================
st.set_page_config(
    page_title="Fleet Compliance Dashboard",
    page_icon="🚛",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ==================================================
# CSS
# ==================================================
st.markdown("""
<style>
:root { --primary-color: #0ea5e9 !important; }
html, body, [class*="css"] { font-family: 'Inter', 'Segoe UI', -apple-system, sans-serif; }
#MainMenu {visibility: hidden;}
footer {visibility: hidden;}
header[data-testid="stHeader"] { background: transparent !important; box-shadow: none !important; }
button[data-testid="baseButton-headerNoPadding"] { visibility: visible !important; opacity: 1 !important; color: #0f172a !important; }
.block-container { padding-top: 2rem !important; padding-bottom: 0.5rem !important; max-width: 100% !important; }
section.main > div { padding-bottom: 0 !important; }
div[data-testid="stAppViewContainer"] > .main { padding-bottom: 0 !important; }
.element-container:has(div[data-testid="stDataFrame"]) { margin-bottom: 0 !important; }
div[data-testid="stDataFrame"] { margin-bottom: 0 !important; }

[data-testid="stMetric"] {
    background: linear-gradient(135deg, #ffffff 0%, #f8fafc 100%);
    border: 1px solid #e2e8f0;
    border-radius: 12px;
    padding: 16px 20px;
    box-shadow: 0 1px 3px rgba(0,0,0,0.04);
    transition: transform 0.15s ease, box-shadow 0.15s ease;
}
[data-testid="stMetric"]:hover { transform: translateY(-2px); box-shadow: 0 4px 12px rgba(0,0,0,0.08); }
[data-testid="stMetricLabel"] { font-size: 13px !important; font-weight: 600; color: #64748b !important; text-transform: uppercase; letter-spacing: 0.5px; }
[data-testid="stMetricValue"] { font-size: 28px !important; font-weight: 700; color: #0f172a !important; }

div[data-testid="stDialog"] > div:first-child {
    border-radius: 16px !important;
    box-shadow: 0 20px 60px rgba(0,0,0,0.15) !important;
    border: 1px solid #e2e8f0 !important;
    background: #ffffff !important;
    padding: 8px 4px !important;
}
div[data-testid="stDialog"] h2 { font-size: 20px !important; font-weight: 700 !important; color: #0f172a !important; padding-bottom: 12px; border-bottom: 1px solid #f1f5f9; margin-bottom: 16px !important; }
div[data-testid="stDialog"] h3 { font-size: 15px !important; font-weight: 700 !important; color: #334155 !important; text-transform: uppercase; letter-spacing: 0.6px; margin-bottom: 16px !important; padding-bottom: 6px; border-bottom: 2px solid #0ea5e9; display: inline-block; }
div[data-testid="stDialog"] input, div[data-testid="stDialog"] textarea {
    background-color: #f8fafc !important;
    border: 1px solid #e2e8f0 !important;
    border-radius: 8px !important;
    color: #0f172a !important;
    font-size: 14px !important;
    font-weight: 500 !important;
    padding: 10px 14px !important;
    -webkit-text-fill-color: #0f172a !important;
    opacity: 1 !important;
}
div[data-testid="stDialog"] label { font-size: 11px !important; font-weight: 600 !important; color: #64748b !important; text-transform: uppercase; letter-spacing: 0.5px; margin-bottom: 4px !important; }
div[data-testid="stDialog"] input:disabled { color: #0f172a !important; -webkit-text-fill-color: #0f172a !important; background-color: #f8fafc !important; opacity: 1 !important; cursor: default !important; }
div[data-testid="stDialog"] .stButton > button, div[data-testid="stDialog"] .stLinkButton > a {
    border-radius: 8px !important;
    font-weight: 600 !important;
    font-size: 13px !important;
    padding: 10px 16px !important;
    border: 1px solid #e2e8f0 !important;
    background: #ffffff !important;
    color: #0f172a !important;
    transition: all 0.15s ease !important;
    text-align: center !important;
    text-decoration: none !important;
}
div[data-testid="stDialog"] .stButton > button:hover, div[data-testid="stDialog"] .stLinkButton > a:hover {
    background: #0ea5e9 !important;
    color: #ffffff !important;
    border-color: #0ea5e9 !important;
    transform: translateY(-1px);
    box-shadow: 0 4px 12px rgba(14,165,233,0.3);
}

div[data-testid="stDataFrame"] { border-radius: 12px; overflow: hidden; border: 1px solid #e2e8f0; }
div[data-testid="stDataFrame"] td:first-child { color: #0ea5e9 !important; font-weight: 600 !important; }
div[data-testid="stDataFrame"] input[type="checkbox"] { accent-color: #0ea5e9 !important; }

section[data-testid="stSidebar"] { background: linear-gradient(180deg, #f8fafc 0%, #ffffff 60%, #f8fafc 100%); border-right: 1px solid #e2e8f0; }
section[data-testid="stSidebar"] h3 { font-size: 18px !important; font-weight: 800 !important; color: #0f172a !important; letter-spacing: -0.4px; margin-top: 8px !important; margin-bottom: 20px !important; padding-bottom: 12px; border-bottom: 2px solid #0ea5e9; display: inline-block; }
section[data-testid="stSidebar"] label { font-size: 11px !important; font-weight: 700 !important; color: #64748b !important; text-transform: uppercase; letter-spacing: 0.8px; margin-bottom: 6px !important; }
section[data-testid="stSidebar"] div[role="radiogroup"] { display: flex; gap: 8px; margin-top: 4px; }
section[data-testid="stSidebar"] div[role="radiogroup"] > label {
    flex: 1;
    background: #ffffff;
    border: 1px solid #e2e8f0;
    border-radius: 10px;
    padding: 10px 8px !important;
    text-align: center;
    cursor: pointer;
    transition: all 0.18s ease;
    font-size: 12px !important;
    font-weight: 700 !important;
    color: #334155 !important;
    text-transform: uppercase !important;
    letter-spacing: 0.5px !important;
    margin: 0 !important;
    white-space: nowrap !important;
    overflow: hidden !important;
}
section[data-testid="stSidebar"] div[role="radiogroup"] > label:hover { border-color: #0ea5e9 !important; background: #f0f9ff !important; color: #0284c7 !important; }
section[data-testid="stSidebar"] div[role="radiogroup"] > label:has(input:checked), section[data-testid="stSidebar"] div[role="radiogroup"] > label[data-checked="true"] {
    background: linear-gradient(135deg, #0ea5e9 0%, #0284c7 100%) !important;
    border-color: #0ea5e9 !important;
    color: #ffffff !important;
    box-shadow: 0 4px 12px rgba(14,165,233,0.3) !important;
}
section[data-testid="stSidebar"] div[role="radiogroup"] > label > div:first-child, section[data-testid="stSidebar"] div[role="radiogroup"] > label > div:first-child * { display: none !important; width: 0 !important; height: 0 !important; margin: 0 !important; }
section[data-testid="stSidebar"] input[type="radio"] { display: none !important; }
section[data-testid="stSidebar"] div[data-baseweb="select"] > div { background: #ffffff !important; border: 1px solid #e2e8f0 !important; border-radius: 10px !important; transition: all 0.18s ease; font-weight: 500; }
section[data-testid="stSidebar"] div[data-baseweb="select"] > div:hover { border-color: #0ea5e9 !important; box-shadow: 0 0 0 3px rgba(14,165,233,0.1); }
section[data-testid="stSidebar"] input[type="text"] { background: #ffffff !important; border: 1px solid #e2e8f0 !important; border-radius: 10px !important; color: #0f172a !important; font-size: 14px !important; padding: 10px 14px !important; -webkit-text-fill-color: #0f172a !important; }
section[data-testid="stSidebar"] input[type="text"]:focus { border-color: #0ea5e9 !important; box-shadow: 0 0 0 3px rgba(14,165,233,0.15) !important; }
section[data-testid="stSidebar"] .stButton > button { background: linear-gradient(135deg, #0ea5e9 0%, #0284c7 100%) !important; border: none !important; color: #ffffff !important; border-radius: 10px !important; font-weight: 700 !important; font-size: 13px !important; padding: 12px 16px !important; letter-spacing: 0.3px; transition: all 0.2s ease !important; box-shadow: 0 4px 12px rgba(14,165,233,0.25); }
section[data-testid="stSidebar"] .stButton > button:hover { transform: translateY(-2px); box-shadow: 0 8px 20px rgba(14,165,233,0.4) !important; background: linear-gradient(135deg, #0284c7 0%, #0369a1 100%) !important; }
section[data-testid="stSidebar"] hr { border-color: #e2e8f0 !important; margin: 20px 0 !important; }
section[data-testid="stSidebar"] img { margin-bottom: 4px; }

.small-updated {
    text-align: left;
    font-size: 11px;
    color: #94a3b8;
    padding: 12px 0 0 0;
    letter-spacing: 0.3px;
}
.small-updated strong { color: #64748b; font-weight: 600; }

h1 { font-size: 28px !important; font-weight: 800 !important; color: #0f172a !important; letter-spacing: -0.5px; }
button[kind="primary"], .stButton > button[kind="primary"] {
    background: linear-gradient(135deg, #0ea5e9 0%, #0284c7 100%) !important;
    border-color: #0ea5e9 !important;
    color: #ffffff !important;
}
</style>
""", unsafe_allow_html=True)


# ==================================================
# CONFIG
# ==================================================
EXCEL_FILE = "data.xlsx"

COMPLIANCE_SHEETS_ATTEMPT = {
    "SWPE": ["SWPE", "swpe", "Swpe", "SWPE ", " SWPE"],
    "SWIN": ["SWIN", "swin", "Swin", "SWIN ", " SWIN"],
}

BASE_DIR  = os.path.dirname(os.path.abspath(__file__))
FILE_PATH = os.path.join(BASE_DIR, EXCEL_FILE)

RC_COLUMN_LETTER = "H"

DOC_CONFIG = {
    "Insurance": {"expiry": "Insurance End Date",  "days": "Insurance Days Left", "col": "Q"},
    "Fitness":   {"expiry": "Fitness Expiry Date", "days": "Fitness Days Left",   "col": "V"},
    "MV Tax":    {"expiry": "MV Tax Expiry Date",  "days": "MV Tax Days Left",    "col": "AA"},
    "Permit":    {"expiry": "Permit Expiry Date",  "days": "Permit Days Left",    "col": "AE"},
    "TP":        {"expiry": "TP Expiry Date",      "days": "TP Days Left",        "col": "AL"},
}

SHAREPOINT_SEARCH_BASE = (
    "https://steelworkspower-my.sharepoint.com/personal/"
    "utkarsh_kashyap_steelworks_in/_layouts/15/onedrive.aspx?q="
)

DUE_DATE_ADVANCE_DAYS = 10


# ==================================================
# SESSION STATE
# ==================================================
if "dialog_open" not in st.session_state:
    st.session_state.dialog_open = False
if "dialog_vehicle" not in st.session_state:
    st.session_state.dialog_vehicle = None
if "search_query" not in st.session_state:
    st.session_state.search_query = ""
if "selected_month" not in st.session_state:
    st.session_state.selected_month = None


# ==================================================
# AUTO-DETECT SHEET NAMES
# ==================================================
@st.cache_data(ttl=60)
def get_actual_sheets():
    if not os.path.exists(FILE_PATH):
        return []
    try:
        with open(FILE_PATH, "rb") as f:
            data = f.read()
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
        names = wb.sheetnames
        wb.close()
        return names
    except Exception:
        return []


def get_last_updated():
    if not os.path.exists(FILE_PATH):
        return None
    mtime = os.path.getmtime(FILE_PATH)
    return datetime.datetime.fromtimestamp(mtime)


# ==================================================
# LOAD HELPERS
# ==================================================
@st.cache_data(ttl=60)
def load_sheet(sheet_name):
    if not os.path.exists(FILE_PATH):
        st.error(f"❌ File not found: {FILE_PATH}")
        st.stop()

    for attempt in range(3):
        try:
            with open(FILE_PATH, "rb") as f:
                data = f.read()
            df = pd.read_excel(io.BytesIO(data), sheet_name=sheet_name, header=0, engine="openpyxl")
            df.columns = df.columns.str.strip()
            return df
        except PermissionError:
            time.sleep(1)
        except Exception as e:
            st.error(f"❌ Error reading sheet '{sheet_name}': {e}")
            st.stop()

    st.error("❌ File is locked. Close Excel and click Rerun.")
    st.stop()


@st.cache_data(ttl=60)
def load_hyperlinks(sheet_name):
    if not os.path.exists(FILE_PATH):
        return {}
    try:
        with open(FILE_PATH, "rb") as f:
            data = f.read()
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
        ws = wb[sheet_name]

        links = {}
        for row in ws.iter_rows():
            for cell in row:
                url = None
                if cell.hyperlink and cell.hyperlink.target:
                    url = cell.hyperlink.target
                elif isinstance(cell.value, str) and cell.value.upper().startswith("=HYPERLINK"):
                    m = re.match(r'=HYPERLINK\(\s*"([^"]+)"', cell.value, re.IGNORECASE)
                    if m:
                        url = m.group(1)
                elif isinstance(cell.value, str) and cell.value.strip().lower().startswith(("http://", "https://")):
                    url = cell.value.strip()
                if url:
                    links.setdefault(cell.row, {})[cell.column_letter] = url
        return links
    except Exception as e:
        st.warning(f"Could not read hyperlinks: {e}")
        return {}


def clean_days_left(series):
    def parse(v):
        if pd.isna(v): return None
        if isinstance(v, (int, float)): return int(v)
        s = str(v).strip().upper()
        if "EXPIRED" in s:       return -1
        if s == "LTT":           return 99999
        if s == "ACTIVE":        return 9999
        if s == "EXPIRING SOON": return 15
        m = re.search(r"-?\d+", s)
        return int(m.group()) if m else None
    return series.apply(parse)


def status_icon(days):
    if days is None or (isinstance(days, float) and pd.isna(days)):
        return "⚪ N/A"
    if not isinstance(days, (int, float)):
        s = str(days).strip().upper()
        if "EXPIRED" in s: return "🔴 EXPIRED"
        if "LTT" in s:     return "🟢 LTT"
        if "ACTIVE" in s:  return "🟢 ACTIVE"
        m = re.search(r"-?\d+", s)
        if not m: return "⚪ N/A"
        days = int(m.group())
    d = int(days)
    if d < 0:   return "🔴 EXPIRED"
    if d <= 30: return "🟡 EXPIRING"
    return "🟢 ACTIVE"


def get_doc_url(doc_label, col_letter, hyperlinks, excel_row, vehicle, veh_col, sheet_name=None):
    url = hyperlinks.get(excel_row, {}).get(col_letter)
    if url:
        return url
    try:
        with open(FILE_PATH, "rb") as f:
            _data = f.read()
        _wb = openpyxl.load_workbook(io.BytesIO(_data), data_only=False)

        sheets_to_check = [sheet_name] if sheet_name else _wb.sheetnames

        for sn in sheets_to_check:
            if sn not in _wb.sheetnames:
                continue
            _ws = _wb[sn]
            cell = _ws[f"{col_letter}{excel_row}"]
            if cell.hyperlink and cell.hyperlink.target:
                return cell.hyperlink.target
            if isinstance(cell.value, str) and cell.value.upper().startswith("=HYPERLINK"):
                m = re.match(r'=HYPERLINK\(\s*"([^"]+)"', cell.value, re.IGNORECASE)
                if m:
                    return m.group(1)
            if isinstance(cell.value, str) and cell.value.strip().lower().startswith(("http://", "https://")):
                return cell.value.strip()
    except Exception:
        pass
    try:
        veh_no = str(vehicle[veh_col]).strip().replace(" ", "%20")
        return f"{SHAREPOINT_SEARCH_BASE}{doc_label}%20{veh_no}"
    except Exception:
        return None


# ==================================================
# MONTHLY SUMMARY
# ==================================================
def build_monthly_summary(df, months_ahead=12, advance_days=DUE_DATE_ADVANCE_DAYS):
    today = pd.Timestamp.today().normalize()
    cutoff = today + pd.DateOffset(months=months_ahead)

    summary = {}

    for doc_name, cfg in DOC_CONFIG.items():
        col = cfg["expiry"]
        if col not in df.columns:
            continue

        for _, row in df.iterrows():
            val = row.get(col)
            if pd.isna(val):
                continue
            try:
                expiry = pd.to_datetime(val)
            except Exception:
                continue

            due = expiry - pd.Timedelta(days=advance_days)
            if due < today or due > cutoff:
                continue

            key = (due.year, due.month)
            if key not in summary:
                summary[key] = {
                    "year": due.year,
                    "month": due.month,
                    "month_name": due.strftime("%B %Y"),
                    "total": 0,
                    "docs": {},
                    "days_left_min": None,
                    "items": [],
                }

            summary[key]["total"] += 1
            summary[key]["docs"][doc_name] = summary[key]["docs"].get(doc_name, 0) + 1

            days_left = (due - today).days
            if summary[key]["days_left_min"] is None or days_left < summary[key]["days_left_min"]:
                summary[key]["days_left_min"] = days_left

            summary[key]["items"].append({
                "Vehicle No": row.get(df.columns[0], ""),
                "Vehicle Name": row.get(df.columns[1], ""),
                "Document": doc_name,
                "Expiry Date": expiry,
                "Due Date": due,
                "Days Until Due": days_left,
            })

    return [summary[k] for k in sorted(summary.keys())]


# ==================================================
# SIDEBAR
# ==================================================
with st.sidebar:
    logo_path = os.path.join(BASE_DIR, "SteelworksLogo.png")
    if os.path.exists(logo_path):
        st.image(logo_path, use_container_width=True)
        st.markdown("<div style='height:12px'></div>", unsafe_allow_html=True)

    st.markdown("### 🔍  Filters")

    search_query = st.text_input("🔎 Search Anything",
                                 value=st.session_state.search_query,
                                 placeholder="Vehicle no., name, chassis, engine, policy...",
                                 key="search_input")
    st.session_state.search_query = search_query

    st.markdown("<div style='height:8px'></div>", unsafe_allow_html=True)

    entity = st.radio("Select Entity", ["SWPE", "SWIN"], horizontal=True, key="entity_filter")

    actual_sheets = get_actual_sheets()

    real_sheet = None
    for candidate in COMPLIANCE_SHEETS_ATTEMPT[entity]:
        for s in actual_sheets:
            if s.strip().upper() == candidate.strip().upper():
                real_sheet = s
                break
        if real_sheet:
            break

    if real_sheet is None:
        st.error(f"❌ Sheet '{entity}' not found in file.")
        st.stop()

    df_all     = load_sheet(real_sheet)
    hyperlinks = load_hyperlinks(real_sheet)

    if search_query.strip():
        q = search_query.strip().lower()
        mask = df_all.apply(
            lambda row: q in " ".join(str(v).lower() for v in row.values if pd.notna(v)),
            axis=1,
        )
        df_all = df_all[mask]

    doc_type = st.selectbox("Document Type", list(DOC_CONFIG.keys()), key="doc_filter")
    cfg = DOC_CONFIG[doc_type]

    if cfg["expiry"] in df_all.columns:
        df_all["_Expiry"]   = pd.to_datetime(df_all[cfg["expiry"]], errors="coerce")
        df_all["_DaysLeft"] = (df_all["_Expiry"] - pd.Timestamp.today()).dt.days
    else:
        df_all["_Expiry"]   = pd.NaT
        df_all["_DaysLeft"] = None

    if df_all["_DaysLeft"].isna().all() and cfg["days"] in df_all.columns:
        df_all["_DaysLeft"] = clean_days_left(df_all[cfg["days"]])

    df_all["_DaysLeft"] = pd.to_numeric(df_all["_DaysLeft"], errors="coerce")
    df_all["_Year"]  = df_all["_Expiry"].dt.year
    df_all["_Month"] = df_all["_Expiry"].dt.month

    years = ["All"] + sorted(df_all["_Year"].dropna().unique().astype(int).tolist())
    year  = st.selectbox("Year", years, key="year_filter")

    month_names = {1:"January", 2:"February", 3:"March", 4:"April", 5:"May", 6:"June",
                   7:"July", 8:"August", 9:"September", 10:"October", 11:"November", 12:"December"}
    month_sel = st.selectbox("Month", ["All"] + list(month_names.values()), key="month_filter")
    month = 0 if month_sel == "All" else [k for k, v in month_names.items() if v == month_sel][0]

    df_filtered = df_all.copy()
    if year != "All":
        df_filtered = df_filtered[df_filtered["_Year"] == year]
    if month != 0:
        df_filtered = df_filtered[df_filtered["_Month"] == month]

    st.divider()

    if st.button("🔄  Reset", use_container_width=True):
        st.cache_data.clear()
        st.session_state.dialog_open    = False
        st.session_state.dialog_vehicle = None
        st.session_state.search_query   = ""
        st.session_state.selected_month = None
        st.rerun()

    _dt = get_last_updated()
    if _dt is not None:
        st.markdown(
            f'<div class="small-updated">Last updated<br>'
            f'<strong>{_dt.strftime("%d %b %Y, %I:%M %p")}</strong></div>',
            unsafe_allow_html=True,
        )


# ==================================================
# HERO HEADER
# ==================================================
st.markdown(f"""
<div style="
    background: linear-gradient(135deg, #1e293b 0%, #0f172a 100%);
    border-radius: 14px;
    padding: 24px 28px;
    color: white;
    margin-bottom: 20px;
    box-shadow: 0 4px 16px rgba(15,23,42,0.25);
">
    <div style="font-size: 12px; font-weight: 600; letter-spacing: 1.5px; opacity: 0.7; text-transform: uppercase; margin-bottom: 6px;">Fleet Compliance</div>
    <div style="font-size: 30px; font-weight: 800; letter-spacing: -0.8px;">{entity} Dashboard</div>
</div>
""", unsafe_allow_html=True)


# ==================================================
# KPI CARDS
# ==================================================
c1, c2, c3, c4 = st.columns(4)
c1.metric("Total Vehicles", len(df_filtered))
c2.metric("🟢 Active",        int((df_filtered["_DaysLeft"] > 30).sum()))
c3.metric("🟡 Expiring Soon", int(((df_filtered["_DaysLeft"] >= 0) & (df_filtered["_DaysLeft"] <= 30)).sum()))
c4.metric("🔴 Expired",       int((df_filtered["_DaysLeft"] < 0).sum()))

st.markdown("<div style='height:24px'></div>", unsafe_allow_html=True)


# ==================================================
# 📅 MONTHLY DUE DATE BREAKDOWN
# ==================================================
st.subheader("📅 Monthly Due Date Breakdown")
st.caption(f"Due date = Expiry Date − {DUE_DATE_ADVANCE_DAYS} days · Click any month to see the vehicle list")

monthly = build_monthly_summary(df_all, months_ahead=12, advance_days=DUE_DATE_ADVANCE_DAYS)

if not monthly:
    st.info("No documents due in the next 12 months.")
else:
    cols_per_row = 3
    for i in range(0, len(monthly), cols_per_row):
        row_items = monthly[i:i + cols_per_row]
        row_cols = st.columns(cols_per_row)

        for col_widget, item in zip(row_cols, row_items):
            with col_widget:
                days = item["days_left_min"] if item["days_left_min"] is not None else 999

                if days < 0: icon = "🔴"
                elif days <= 30: icon = "🟡"
                else: icon = "🟢"

                pills_text = " · ".join(f"{count} {doc}" for doc, count in item["docs"].items())
                btn_label = f"{icon}  {item['month_name']}   ({item['total']})\n{pills_text}"

                key = f"month_{item['year']}_{item['month']}"
                if st.button(btn_label, key=key, use_container_width=True):
                    st.session_state.selected_month = (item["year"], item["month"])
                    st.rerun()

    # ---------- Drill-down ----------
    if st.session_state.selected_month is not None:
        sel_year, sel_month = st.session_state.selected_month
        sel_data = next((m for m in monthly if m["year"] == sel_year and m["month"] == sel_month), None)

        if sel_data:
            st.markdown("<div style='height:16px'></div>", unsafe_allow_html=True)
            st.markdown(f"### 📋 Vehicles due in {sel_data['month_name']}")

            colA, colB = st.columns([4, 1])
            with colB:
                if st.button("✖  Close", key="close_month", use_container_width=True):
                    st.session_state.selected_month = None
                    st.rerun()

            items_df = pd.DataFrame(sel_data["items"])
            if not items_df.empty:
                items_df["Expiry Date"] = pd.to_datetime(items_df["Expiry Date"]).dt.strftime("%d-%b-%Y")
                items_df["Due Date"]    = pd.to_datetime(items_df["Due Date"]).dt.strftime("%d-%b-%Y")

                items_df = items_df[["Vehicle No", "Vehicle Name", "Document",
                                     "Expiry Date", "Due Date", "Days Until Due"]]

                st.dataframe(
                    items_df,
                    use_container_width=True,
                    hide_index=True,
                    column_config={
                        "Days Until Due": st.column_config.NumberColumn("Days Until Due", format="%d"),
                    },
                )

st.markdown("<div style='height:24px'></div>", unsafe_allow_html=True)
st.divider()


# ==================================================
# MAIN TABLE — Insurance Records / Fitness Records etc.
# ==================================================
st.subheader(f"📋 {doc_type} Records")

if len(df_filtered) == 0:
    st.warning("⚠️ No vehicles match the current search or filters.")
    st.stop()

VEH_COL  = df_filtered.columns[0]
NAME_COL = df_filtered.columns[1]

display_cols = [VEH_COL, NAME_COL, cfg["expiry"], "_DaysLeft"]
available    = [c for c in display_cols if c in df_filtered.columns]

display_df = df_filtered[available].copy()
display_df.columns = ["Vehicle No", "Vehicle Name", "Expiry Date", "Days Left"][:len(available)]
display_df["Days Left"] = pd.to_numeric(display_df["Days Left"], errors="coerce")
display_df["Status"] = display_df["Days Left"].apply(status_icon)

# Header
header_cols = st.columns([2.2, 3.5, 2.2, 1.5, 1.8, 1.2])
header_cols[0].markdown("**Vehicle No**")
header_cols[1].markdown("**Vehicle Name**")
header_cols[2].markdown("**Expiry Date**")
header_cols[3].markdown("**Days Left**")
header_cols[4].markdown("**Status**")
header_cols[5].markdown("**Action**")

st.markdown("<hr style='margin:4px 0; border-color:#e2e8f0;'>", unsafe_allow_html=True)

# Rows
for i, row in display_df.iterrows():
    cols = st.columns([2.2, 3.5, 2.2, 1.5, 1.8, 1.2])

    cols[0].write(row["Vehicle No"])
    cols[1].write(row["Vehicle Name"])

    exp_val = row["Expiry Date"]
    if pd.notna(exp_val):
        try:
            cols[2].write(pd.to_datetime(exp_val).strftime("%d-%b-%Y"))
        except Exception:
            cols[2].write(str(exp_val))
    else:
        cols[2].write("—")

    dl = row["Days Left"]
    cols[3].write(int(dl) if pd.notna(dl) else "—")
    cols[4].write(row["Status"])

    with cols[5]:
        btn_key = f"view_{entity}_{i}_{row['Vehicle No']}"
        if st.button("🔍 View", key=btn_key, use_container_width=True):
            st.session_state.dialog_vehicle = row["Vehicle No"]
            st.session_state.dialog_open    = True
            st.rerun()

    st.markdown("<hr style='margin:2px 0; border-color:#f1f5f9;'>", unsafe_allow_html=True)


# ==================================================
# VEHICLE DETAILS DIALOG
# ==================================================
if st.session_state.dialog_open and st.session_state.dialog_vehicle is not None:
    veh_no_clicked = str(st.session_state.dialog_vehicle).strip().upper()

    match_rows = df_filtered[
        df_filtered[VEH_COL].astype(str).str.strip().str.upper() == veh_no_clicked
    ]

    if match_rows.empty:
        st.session_state.dialog_open    = False
        st.session_state.dialog_vehicle = None
        st.rerun()

    vehicle = match_rows.iloc[0]

    # Find excel row
    excel_row = None
    try:
        with open(FILE_PATH, "rb") as f:
            _data = f.read()
        _wb = openpyxl.load_workbook(io.BytesIO(_data), data_only=True)
        _ws = _wb[real_sheet]
        for r in range(2, _ws.max_row + 1):
            cell_val = _ws.cell(row=r, column=1).value
            if cell_val and str(cell_val).strip().upper() == veh_no_clicked:
                excel_row = r
                break
    except Exception:
        pass

    if excel_row is None:
        excel_row = 2

    @st.dialog(f"🚛  {vehicle.get(NAME_COL, '')}  •  {vehicle[VEH_COL]}", width="large")
    def show_vehicle_details():

        st.markdown(f"""
        <div style="
            background: linear-gradient(135deg, #0ea5e9 0%, #0284c7 100%);
            border-radius: 12px;
            padding: 20px 24px;
            margin-bottom: 24px;
            color: white;
            box-shadow: 0 4px 14px rgba(14,165,233,0.3);
        ">
            <div style="font-size: 11px; font-weight: 600; opacity: 0.85; text-transform: uppercase; letter-spacing: 1.2px; margin-bottom: 4px;">Vehicle</div>
            <div style="font-size: 24px; font-weight: 700; letter-spacing: -0.5px;">{vehicle[VEH_COL]}</div>
            <div style="font-size: 14px; opacity: 0.95; margin-top: 4px;">{vehicle.get(NAME_COL, '')}</div>
        </div>
        """, unsafe_allow_html=True)

        left, right = st.columns([1.7, 1])

        with left:
            st.markdown("### 📝 Vehicle Details")
            info = [
                ("RTO Location", "RTO Location"),
                ("Model No.", "Model No."),
                ("Chassis No.", "Chassis No."),
                ("Engine No", "Engine No"),
                ("Model Year", "Model Year"),
                ("Insurance Company", "Insurance Company"),
                ("Policy No.", "Policy No."),
                ("Insured Value (IDV)", "Insured Value (IDV)"),
                ("Insurance Amount", "Insurance Amount"),
                ("Fully Compliant Vehicle", "Fully Compliant Vehicle"),
            ]
            c1, c2 = st.columns(2)
            for i, (label, col) in enumerate(info):
                if col in vehicle.index and pd.notna(vehicle[col]):
                    val = vehicle[col]
                    if isinstance(val, float):
                        if col in ("Insured Value (IDV)", "Insurance Amount"):
                            val = f"₹ {val:,.0f}"
                        else:
                            val = f"{val:g}"
                    with (c1 if i % 2 == 0 else c2):
                        st.text_input(label, str(val), disabled=True, key=f"fld_{i}_{veh_no_clicked}")

        with right:
            st.markdown("### 📄 Documents")

            rc_url = get_doc_url("RC", RC_COLUMN_LETTER, hyperlinks, excel_row, vehicle, VEH_COL, real_sheet)
            if rc_url:
                st.link_button("📋   View RC", rc_url, use_container_width=True)
            else:
                st.button("❌   RC — No file", disabled=True,
                          use_container_width=True, key=f"nodoc_rc_{veh_no_clicked}")

            for doc_name, doc_cfg in DOC_CONFIG.items():
                doc_url = get_doc_url(doc_name, doc_cfg["col"], hyperlinks, excel_row, vehicle, VEH_COL, real_sheet)
                if doc_url:
                    st.link_button(f"📄   View {doc_name}", doc_url, use_container_width=True)
                else:
                    st.button(f"❌   {doc_name} — No file",
                              disabled=True, use_container_width=True,
                              key=f"nodoc_{doc_name}_{veh_no_clicked}")

        st.markdown("""
        <div style="text-align:center; color:#94a3b8; font-size:11px; padding:16px 0 4px 0; letter-spacing:0.3px;">
            Click outside or press <b>ESC</b> to close
        </div>
        """, unsafe_allow_html=True)

    show_vehicle_details()


# ==================================================
# EXPIRING SOON
# ==================================================
st.markdown("<div style='height:20px'></div>", unsafe_allow_html=True)
st.divider()
st.subheader("⚠️ Expiring in 30 Days")

soon = df_filtered[(df_filtered["_DaysLeft"] >= 0) & (df_filtered["_DaysLeft"] <= 30)]

if soon.empty:
    st.success("✅ No documents expiring in the next 30 days.")
else:
    soon_cols  = [VEH_COL, NAME_COL, "_DaysLeft"]
    soon_avail = [c for c in soon_cols if c in soon.columns]
    soon_display = soon[soon_avail].copy()
    soon_display.columns = ["Vehicle No", "Vehicle Name", "Days Left"][:len(soon_avail)]
    st.dataframe(soon_display, use_container_width=True, hide_index=True)


# ==================================================
# FOOTER
# ==================================================
st.markdown(f"""
<div style="text-align:center; color:#94a3b8; font-size:12px; padding:16px 0 8px 0; margin-top:12px; border-top:1px solid #f1f5f9;">
    © {pd.Timestamp.now().year} Steelworks • Fleet Compliance System
</div>
""", unsafe_allow_html=True)